"""
Load true-positive GO term keywords from Supplementary Table 8.

The paper defines cell-type-specific keywords that are used to identify
"true positive" GO terms from GSEA results. These keywords were manually
curated by medical and scientific staff.

Suppl Table 8 structure (from paper's GitHub):
  Columns: Cardiomyocytes, Endothelial, Fibroblasts, Pericytes,
           Immune Cells, Smooth Muscle, Neuronal
  Rows: keyword lists per cell type
"""
import os
import openpyxl
from typing import Dict, List, Set


# Cell type name mapping (paper's column names -> our class names)
CELLTYPE_GO_MAPPING = {
    'Cardiomyocytes': 'Cardiomyocytes',
    'Endothelial': 'Endothelial',
    'Fibroblasts': 'Fibroblasts',
    'Pericytes': 'Pericytes',
    'Immune Cells': 'Immune.cells',
    'Smooth Muscle': 'Smooth.Muscle',
    'Neuronal': 'Neuro',
}


def load_go_keywords(xlsx_path: str) -> Dict[str, List[str]]:
    """
    Load GO term keywords from Supplementary Table 8.

    Args:
        xlsx_path: path to Suppl_Table_8_TP_GO_ID_per_cell_type.xlsx

    Returns:
        dict mapping cell type name -> list of keyword strings
        e.g. {'Cardiomyocytes': ['heart', 'cardiac', 'muscle', ...], ...}
    """
    if not os.path.exists(xlsx_path):
        print(f"Warning: GO terms file not found at {xlsx_path}")
        print("  Using built-in defaults from paper Supplementary Table 8.")
        return get_default_go_keywords()

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    # Read header (cell types)
    cell_types = []
    for c in range(1, ws.max_column + 1):
        cell_types.append(ws.cell(1, c).value)

    # Read keywords per cell type
    keywords: Dict[str, List[str]] = {}
    for c, ct in enumerate(cell_types):
        if ct is None:
            continue
        ct_str = str(ct).strip()
        ct_keywords = []
        for r in range(2, ws.max_row + 1):
            val = ws.cell(r, c + 1).value
            if val is not None and str(val).strip():
                ct_keywords.append(str(val).strip().lower())
        # Map to our class name
        our_name = CELLTYPE_GO_MAPPING.get(ct_str, ct_str)
        keywords[our_name] = ct_keywords

    return keywords


def get_default_go_keywords() -> Dict[str, List[str]]:
    """
    Default GO keywords from paper Supplementary Table 8.
    Used if the xlsx file is not available.
    """
    return {
        'Cardiomyocytes': [
            'heart', 'cardiac', 'muscle', 'myocardium', 'contraction',
            'sarcomere', 'cardiogenesis', 'hypertrophy', 'electrophysiology',
            'ion transport', 'gap junction', 'mitochondria', 'angiogenesis',
            'apoptosis', 'fibrosis', 'calcium ion',
        ],
        'Endothelial': [
            'endothelium', 'endothelial', 'heart', 'angiogenesis',
            'vascularization', 'inflammation', 'vascular permeability',
            'platelet activation', 'shear stress', 'cell adhesion',
            'smooth muscle contraction', 'endothelial dysfunction',
            'nitric oxide production', 'extracellular matrix',
            'hypoxia', 'vascular remodeling',
        ],
        'Fibroblasts': [
            'fibroblast', 'extracellular matrix', 'extracellular structure',
            'wound healing', 'wounding', 'tissue repair', 'heart',
            'connective tissue', 'collagen', 'mesenchymal', 'fibronectin',
        ],
        'Pericytes': [
            'pericyte', 'blood vessel morphogenesis', 'heart',
            'regulation of vascular permeability', 'angiogenesis',
            'vascular smooth muscle cell development', 'actin filament',
            'basement membrane organization', 'mesenchymal',
            'extracellular matrix',
        ],
        'Immune.cells': [
            'immune response', 'heart', 'inflammation', 'leukocyte activation',
            'cytokine production', 'antigen presentation', 't cell activation',
            'b cell differentiation', 'immune cell migration',
            'macrophage activation', 'immune system development',
            'interleukin signaling',
        ],
        'Smooth.Muscle': [
            'muscle contraction', 'heart', 'smooth muscle',
            'actin filament', 'muscle development', 'vascular smooth muscle',
            'vasoconstriction', 'extracellular matrix organization',
            'vascular remodeling', 'response to mechanical stimulus',
            'muscle hypertrophy',
        ],
        'Neuro': [
            'neuron differentiation', 'axon', 'neurogenesis', 'heart',
            'neuronal', 'synapse organization', 'neurotransmitter release',
            'dendrite development', 'axonogenesis', 'neuron',
            'nerve growth factor signaling', 'membrane potential',
        ],
    }


def match_go_terms_to_keywords(go_terms: List[str],
                                keywords: List[str]) -> Set[str]:
    """
    Check which GO terms match any of the true-positive keywords.

    A GO term is a "true positive" if its name/description contains
    any of the keywords (case-insensitive substring match).

    Args:
        go_terms: list of GO term names/descriptions
        keywords: list of keyword strings to match against

    Returns:
        set of GO term names that match at least one keyword
    """
    matched = set()
    keywords_lower = [k.lower() for k in keywords]
    for term in go_terms:
        term_lower = term.lower()
        for kw in keywords_lower:
            if kw in term_lower:
                matched.add(term)
                break
    return matched


def compute_gsea_f1(found_terms: Set[str], true_positive_terms: Set[str],
                    all_possible_terms: Set[str]) -> Dict[str, float]:
    """
    Compute F1-score for GSEA results against true positive GO terms.

    Following the paper's approach:
      - True positives: found terms that are in the true positive set
      - False positives: found terms that are NOT in the true positive set
      - False negatives: true positive terms that were NOT found
      - True negatives: all other terms (not found, not true positive)

    Args:
        found_terms: GO terms identified as significant by GSEA
        true_positive_terms: GO terms matching the keyword set
        all_possible_terms: all GO terms in the gene set database

    Returns:
        dict with precision, recall, f1, tp, fp, fn, tn counts
    """
    tp = len(found_terms & true_positive_terms)
    fp = len(found_terms - true_positive_terms)
    fn = len(true_positive_terms - found_terms)
    tn = len(all_possible_terms - found_terms - true_positive_terms)

    precision = tp / (tp + fp) if (tp + fp) > 0 else 0.0
    recall = tp / (tp + fn) if (tp + fn) > 0 else 0.0
    f1 = 2 * precision * recall / (precision + recall) if (precision + recall) > 0 else 0.0

    return {
        'precision': precision,
        'recall': recall,
        'f1': f1,
        'tp': tp,
        'fp': fp,
        'fn': fn,
        'tn': tn,
    }
